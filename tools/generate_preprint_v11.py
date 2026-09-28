#!/usr/bin/env python3
"""Generate preprint Version 1.1 (outreach revision). No new experimental results."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Emu, Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures_v11"
OUT = ROOT / "Structural_Identifiability_Epistemic_Value_Active_Inference_v1.1.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE = "1F3A5F"
BOX_FILL = "EEF3F8"
HEAD_FILL = "1F3A5F"
ALT_FILL = "F4F7FA"
PAGE_W = Inches(8.5)
CONTENT_W = Inches(6.5)  # 8.5 - 1 - 1


def set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color=RULE, sz="8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_width(cell, width_dxa: int):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_bottom_border(paragraph, color=RULE, sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, space_before, space_after, bold in (
        ("Heading 1", 14, 18, 8, True),
        ("Heading 2", 12, 14, 6, True),
        ("Heading 3", 11, 10, 4, True),
    ):
        st = styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Caption" not in [s.name for s in styles]:
        cap = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.color.rgb = MUTED
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(
    doc,
    text,
    *,
    first_indent=True,
    italic=False,
    bold=False,
    center=False,
    size=11,
    space_after=8,
    color=BODY,
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(18) if first_indent else Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def mixed_para(doc, spans, *, first_indent=True, center=False, space_after=8, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(18) if first_indent else Pt(0)
    for text, kwargs in spans:
        run = p.add_run(text)
        set_run_font(
            run,
            size=kwargs.get("size", size),
            bold=kwargs.get("bold", False),
            italic=kwargs.get("italic", False),
            color=kwargs.get("color", BODY),
        )
    return p


def equation(doc, text, tag=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run(text)
    set_run_font(run, name="Cambria Math", size=12, italic=True, color=NAVY)
    if tag:
        tab = p.add_run(f"    ({tag})")
        set_run_font(tab, size=10, italic=False, color=MUTED)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.35)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if isinstance(item, str):
            run = p.add_run(item)
            set_run_font(run, size=11)
        else:
            for text, kwargs in item:
                run = p.add_run(text)
                set_run_font(
                    run, size=11, bold=kwargs.get("bold", False), italic=kwargs.get("italic", False)
                )


def callout(doc, title, body, fill=BOX_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = 9360  # 6.5 inches in DXA
    table.columns[0].width = width
    cell = table.cell(0, 0)
    set_cell_width(cell, width)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=RULE, sz="12")
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_before = Pt(6)
    p1.paragraph_format.space_after = Pt(2)
    r = p1.add_run(title)
    set_run_font(r, size=10, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(6)
    r2 = p2.add_run(body)
    set_run_font(r2, size=11, italic=False, color=BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_table(doc, headers, rows, col_widths=None, caption=None):
    ncols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if col_widths is None:
        col_widths = [int(9360 / ncols)] * ncols
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_width(cell, col_widths[j])
        shade_cell(cell, HEAD_FILL)
        set_cell_borders(cell, color=HEAD_FILL, sz="4")
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=9, bold=True, color=RGBColor(255, 255, 255))
    for i, row in enumerate(rows):
        prevent_row_split(table.rows[i + 1])
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            set_cell_width(cell, col_widths[j])
            shade_cell(cell, ALT_FILL if i % 2 else "FFFFFF")
            set_cell_borders(cell, color="C5CDD6", sz="4")
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            set_run_font(r, size=9, color=BODY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(12)
        r = cap.add_run(caption)
        set_run_font(r, size=10, italic=True, color=MUTED)
    return table


def add_figure(doc, path: Path, caption: str, width=Inches(6.0)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run_font(r, size=10, italic=True, color=MUTED)


def setup_section(doc: Document):
    section = doc.sections[0]
    section.page_width = PAGE_W
    section.page_height = Inches(11)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Structural identifiability of epistemic value  ·  Version 1.1")
    set_run_font(hr, size=9, italic=True, color=MUTED)
    add_bottom_border(hp, color="C5CDD6", sz="6")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("METHOD PROPOSED  ·  SYNTHETICALLY IDENTIFIABLE  ·  HUMAN VALIDATION PENDING")
    set_run_font(fr, size=8, color=MUTED)
    fp.add_run("    |    ")
    # page number field
    run = fp.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    run._r.append(fld1)
    run2 = fp.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._r.append(instr)
    run3 = fp.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run3._r.append(fld2)
    for r in (run, run2, run3):
        set_run_font(r, size=8, color=MUTED)
    fp.add_run("    |    Internal ID: AIF-KILL-001")
    set_run_font(fp.runs[-1], size=8, color=MUTED)


def regenerate_figures():
    FIG.mkdir(exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )

    # Figure 1 — adversarial sequence
    fig, ax = plt.subplots(figsize=(8.2, 2.4))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2)
    ax.axis("off")
    boxes = [
        (0.4, "Reward-only\nv1–v3"),
        (2.9, "Generic curiosity\nv4–v5"),
        (5.4, "Exact-MI target\nv6"),
        (7.9, "Same-objective\nceiling v7"),
    ]
    for x, label in boxes:
        rect = plt.Rectangle((x, 0.7), 1.8, 0.9, fill=False, linewidth=1.2, edgecolor="#1F3A5F")
        ax.add_patch(rect)
        ax.text(x + 0.9, 1.15, label, ha="center", va="center", fontsize=9, color="#1F3A5F")
    for x0, x1 in ((2.2, 2.9), (4.7, 5.4), (7.2, 7.9)):
        ax.annotate(
            "",
            xy=(x1, 1.15),
            xytext=(x0, 1.15),
            arrowprops=dict(arrowstyle="->", color="#1F3A5F", lw=1.2),
        )
    ax.text(
        5,
        0.25,
        "Increasing adversarial strength  →  narrower licensed conclusion",
        ha="center",
        fontsize=10,
        color="#333333",
    )
    fig.tight_layout()
    fig.savefig(FIG / "fig1_sequence.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    # Figure 3 — v4 JS from reported numbers
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    xs = ["Median min JS", "90th percentile min JS", "Maximum min JS"]
    ys = [5.5e-7, 5.9e-6, 2.0e-5]
    ax.bar(xs, ys, color="#2F5D8A", width=0.55)
    ax.ticklabel_format(axis="y", style="sci", scilimits=(0, 0))
    ax.set_ylabel("Minimum Jensen–Shannon divergence")
    ax.set_title("v4: generic curiosity nearly reproduces exact-MI choice distributions")
    fig.tight_layout()
    fig.savefig(FIG / "fig3_v4.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    # Figure 4 — four candidate pairs from the retained table
    fig, ax = plt.subplots(figsize=(6.6, 3.8))
    pairs = np.arange(1, 5)
    aif = [0.277, 0.270, 0.259, 0.245]
    cur = [0.0, 2.22e-16, 0.0, 0.0]
    w = 0.36
    ax.bar(
        pairs - w / 2, aif, width=w, label="Exact-MI / reduced AIF  |ΔP(sample)|", color="#2F5D8A"
    )
    ax.bar(
        pairs + w / 2, cur, width=w, label="Accuracy-based curiosity  |ΔP(sample)|", color="#D4A017"
    )
    ax.set_xticks(pairs)
    ax.set_xlabel("Candidate matched-channel pair")
    ax.set_ylabel("Absolute within-pair sampling contrast")
    ax.set_title("v6: same prior, mean accuracy and instrumental value; different MI")
    ax.legend(frameon=False, fontsize=8)
    ax.set_ylim(0, 0.32)
    fig.tight_layout()
    fig.savefig(FIG / "fig4_v6_contrast.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    # Figure 5 — recovery table
    fig, ax = plt.subplots(figsize=(6.6, 3.8))
    trials = [60, 120, 240, 480]
    aif_r = [0.972, 0.961, 0.983, 0.994]
    cur_r = [0.800, 0.896, 0.967, 1.000]
    ax.plot(trials, aif_r, "o-", color="#2F5D8A", label="True exact-MI / reduced AIF family")
    ax.plot(trials, cur_r, "o-", color="#D4A017", label="True coarse-curiosity family")
    ax.axhline(0.80, color="#888888", ls="--", lw=1)
    ax.set_xlabel("Trials per synthetic subject")
    ax.set_ylabel("Correct family recovery")
    ax.set_title("v6: matched asymmetric channels restore family discrimination")
    ax.set_ylim(0, 1.05)
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(FIG / "fig5_v6_recovery.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def build():
    regenerate_figures()
    doc = Document()
    configure_styles(doc)
    setup_section(doc)

    # Title block
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(4)
    r = t.add_run("Structural Identifiability of Epistemic Value\nin Active Inference")
    set_run_font(r, size=20, bold=True, color=NAVY)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(10)
    r = st.add_run("Adversarial Model Comparison and a Matched-Channel Experimental Design")
    set_run_font(r, size=13, italic=True, color=NAVY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    r = meta.add_run("Research note  ·  Version 1.1  ·  28 September 2026")
    set_run_font(r, size=11, color=MUTED)

    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_after = Pt(2)
    r = status.add_run(
        "METHOD PROPOSED  +  SYNTHETICALLY IDENTIFIABLE  +  HUMAN VALIDATION PENDING"
    )
    set_run_font(r, size=10, bold=True, color=NAVY)

    ident = doc.add_paragraph()
    ident.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ident.paragraph_format.space_after = Pt(14)
    add_bottom_border(ident, color=RULE, sz="12")
    r = ident.add_run("Internal project identifier: AIF-KILL-001")
    set_run_font(r, size=10, italic=True, color=MUTED)

    # Abstract
    heading(doc, "Abstract", 1)
    para(
        doc,
        "Active Inference (AIF) assigns policies both pragmatic/extrinsic and epistemic/intrinsic value. "
        "An empirical difficulty is that information-seeking choices can be generated by many non-AIF "
        "mechanisms, so successful model fit need not imply theory-specific support. We report an "
        "adversarial sequence of synthetic tests designed to determine what choice behavior can and "
        "cannot identify. A reduced one-step policy-ranking model is related here to the expected-free-energy "
        "(EFE) decomposition under stated assumptions. Reward-only and flexible Bayes-adaptive rivals were "
        "separable for sufficiently strong epistemic effects, but a generic curiosity bonus rendered ordinary "
        "symmetric-cue tasks nearly observationally equivalent. A search over 175,821 condition pairs did not "
        "yield a fully robust sign-reversal discriminator against the tested curiosity family. We therefore "
        "propose a matched asymmetric-channel design: pairs with equal prior, equal mean cue accuracy, and "
        "equal instrumental sampling value, but different exact mutual information. In the tested synthetic "
        "model class, this design generated substantial exact-MI sampling contrasts while the mean-accuracy "
        "curiosity rival generated approximately none, and model-family recovery exceeded 96% for both families "
        "by 240 trials. However, an alternative controller implementing the same exact mutual-information "
        "objective and policy mapping is behaviorally indistinguishable from the reduced AIF controller. "
        "The proposed human study therefore targets exact-MI-sensitive epistemic valuation relative to specified "
        "alternatives. A close human OED precedent already exists (Nelson et al., 2010). The remaining question "
        "is what an Active Inference framing adds beyond that contrast. The design is not a test of the entirety "
        "of Active Inference, and it cannot uniquely verify the universal Free Energy Principle.",
        first_indent=False,
    )
    mixed_para(
        doc,
        [
            ("Keywords: ", {"bold": True, "italic": True}),
            (
                "Active Inference; expected free energy; epistemic value; mutual information; "
                "identifiability; model comparison; curiosity; experimental design; preregistration.",
                {"italic": True},
            ),
        ],
        first_indent=False,
        space_after=12,
        size=10,
    )

    heading(doc, "1. Introduction", 1)
    heading(doc, "1.1 Motivation: from a general principle to an identifiable prediction", 2)
    para(
        doc,
        "The Free Energy Principle, taken as a general statement about self-organizing systems, is too coarse "
        "to be the object of this study. Scientific content appears only when a specific expected-free-energy "
        "/ Active Inference instantiation yields an identifiable prediction that a strong rival can fail. The "
        "intended inference ladder is therefore",
    )
    equation(
        doc,
        "general FEP  →  specific EFE/AIF model  →  identifiable prediction  →  strong rival",
        None,
    )
    para(
        doc,
        "Compatibility with exploratory data is not confirmation of the principle. A 2026 review of the "
        "empirical science of Active Inference makes the same demand: theory-driven experiments that "
        "adjudicate between accounts, rather than treating fit as support for the framework as a whole [7]. "
        "The present note occupies only the third and fourth rungs of that ladder, and only for epistemic value.",
    )
    heading(doc, "1.2 The identifiability problem for epistemic value", 2)
    para(
        doc,
        "Active Inference proposes that policies are selected by minimizing expected free energy, a quantity "
        "that can be expressed in terms of pragmatic/extrinsic and epistemic/intrinsic value. In standard "
        "discrete-state formulations, epistemic value corresponds to expected information gain about hidden "
        "states or parameters, linking Active Inference to optimal Bayesian design, active learning, and "
        "broader information-seeking frameworks [1,2].",
    )
    para(
        doc,
        "This creates an identifiability problem. Observing exploration, uncertainty reduction, or even "
        "behavior consistent with mutual information does not uniquely identify the theoretical framework "
        "that produced it. Intrinsic-motivation and curiosity models predate much of the contemporary "
        "Active Inference literature [3]. Information-directed sampling uses mutual information explicitly "
        "to balance exploration and exploitation [4]. Human information-gathering research distinguishes "
        "competing norms such as information gain, probability gain, diagnosticity, and impact [5]. "
        "The relevant competitor class is therefore not merely reward-based reinforcement learning, but "
        "alternative information-seeking objectives.",
    )
    para(
        doc,
        "Recent work reinforces the same discipline. A 2026 complexity-matched comparison of AIF and "
        "reinforcement-learning models found overlapping but partially distinct explanations and similar "
        "predictive accuracy [6]. That overlap is expected if a general principle is too coarse and only "
        "specific instantiations can be tested.",
    )
    para(
        doc,
        "The present note develops a candidate discriminative design and, equally importantly, specifies what "
        "that design cannot establish. We do not claim a first demonstration, a first structural discriminator, "
        "or a theoretical breakthrough. We propose a matched asymmetric-channel design and report an adversarial "
        "synthetic sequence that delimits its interpretive reach. All numerical results reported here are "
        "synthetic. No human empirical data have been collected.",
    )
    callout(
        doc,
        "Scope of this note",
        "This paper is one kill-test of epistemic-value identifiability. It is not a test of the Free Energy Principle as a general statement, not a closed-loop causal intervention experiment, not a coarse-graining / scale-gap result, and not clinical evidence. A later positive human contrast would still sit on the third and fourth rungs of the ladder in Section 1.1.",
    )

    heading(doc, "2. The identifiability problem", 1)
    para(
        doc,
        "The project began from a strong empirical ambition:",
        first_indent=False,
    )
    callout(
        doc,
        "Initial claim C0 — not treated as a premise",
        "AIF produces behaviorally distinctive epistemic choices that cannot be reproduced by strong non-AIF alternatives.",
    )
    para(
        doc,
        "C0 was not treated as a premise. Each stage introduced a stronger rival or a stricter identifiability "
        "test. An earlier PASS was allowed to lose interpretive force when a later rival reproduced the same "
        "observables. The resulting discipline is that a positive fit licenses only the weakest conclusion "
        "still standing after the strongest implemented rival has been confronted.",
    )
    add_table(
        doc,
        ["Result", "Maximum licensed conclusion"],
        [
            [
                "Exact-MI model beats reward-only control",
                "Information value improves prediction relative to that rival.",
            ],
            [
                "Exact-MI model beats a flexible reward-based Bayes-adaptive rival",
                "The effect is not absorbed by the tested reward/uncertainty nuisance parameters.",
            ],
            [
                "Generic curiosity matches the exact-MI model",
                "Information seeking is not mechanism-specific.",
            ],
            [
                "Matched-channel exact-MI model beats tested coarse curiosity",
                "Behavior is sensitive to information geometry beyond mean accuracy in the tested model class.",
            ],
            [
                "Another controller uses the same exact-MI objective and policy mapping",
                "Behavior cannot identify the theoretical derivation.",
            ],
        ],
        col_widths=[4200, 5160],
        caption="Table 1. Allowed versus forbidden inferences after the adversarial sequence.",
    )
    para(
        doc,
        "The remainder of the note follows this discipline. Version 1.1 adds an explicit EFE derivation for "
        "the reduced controller, recasts the v7 equivalence as an interpretation boundary, downgrades the "
        "human protocol from “ready” to a candidate, and expands the rival set that a later confirmatory "
        "study must confront.",
    )

    heading(doc, "3. Relation of the reduced model to expected free energy", 1)
    para(
        doc,
        "The main technical revision in Version 1.1 is to state how the reduced controller relates to expected "
        "free energy, and where that relation stops. A reviewer can otherwise say, correctly, that a generic "
        "information-gain controller was compared rather than Active Inference as such.",
    )

    heading(doc, "3.1 General EFE decomposition", 2)
    para(
        doc,
        "For a future policy π, a standard discrete-state formulation expresses expected free energy as "
        "containing an extrinsic/preference term and an epistemic expected-information-gain term, up to an "
        "expected-evidence-bound (ambiguity) term [1,2,10]. Under idealized model alignment, the policy ranking "
        "can be written schematically as",
    )
    equation(doc, "−G(π)  ≈  E_Q[log P_C(o) | π]  +  I_Q(S; O | π)  +  constant", "1")
    para(
        doc,
        "Here P_C(o) encodes prior preferences over outcomes and I_Q(S;O|π) is the expected mutual information "
        "between hidden state and future observation under the policy. Equation (1) is the starting point of "
        "the reduction, not an additional assumption that information gain is unique to Active Inference.",
    )

    heading(doc, "3.2 When the pragmatic term becomes instrumental value", 2)
    para(
        doc,
        "Assume outcome preferences take the exponential form P_C(o) ∝ exp(κ r(o)), where r(o) is the task "
        "reward and κ > 0 sets preference precision. Then",
    )
    equation(doc, "E_Q[log P_C(o) | π]  =  κ · E_Q[r(o) | π]  −  log Z", "2")
    para(
        doc,
        "Because −log Z is policy-independent, ranking policies by −G becomes equivalent to ranking",
    )
    equation(doc, "E_Q[r(o) | π]  +  (1/κ) · I_Q(S; O | π)", "3")
    para(
        doc,
        "After rescaling, the reduced score used in the simulations is therefore",
    )
    equation(doc, "V_red(π)  =  V_instr(π)  +  α_epi · I(S; O | π)", "4")
    para(
        doc,
        "In the one-step task, only SAMPLE has a nonzero epistemic term. The other actions (SAFE, COMMIT_L, "
        "COMMIT_R) are scored by instrumental value alone. This is the reduced Active-Inference controller "
        "used throughout v1–v8. It is a model of an EFE-derived epistemic-value component, not a process-level "
        "implementation of full Active Inference.",
    )

    heading(doc, "3.3 Status of α_epi", 2)
    para(
        doc,
        "The parameter α_epi must not be conflated with softmax temperature. Decision precision β maps scores "
        "to choice probabilities and is a property of the policy mapping, not of EFE. Under the derivation "
        "above, α_epi = 1/κ is the inverse preference precision: a derived relative scale between reward "
        "units and nats of expected information gain. In that reading it is not a phenomenological curiosity "
        "slider added on top of Active Inference.",
    )
    para(
        doc,
        "In the reported simulations, however, the reward scale is held fixed and α_epi is varied as a free "
        "parameter. That usage is a phenomenological extension of the reduced ranking model. It is useful for "
        "recovery and identifiability analysis, but it is not a claim that a freely weighted information-gain "
        "term is the unique or canonical parameterization of full Active Inference.",
    )

    heading(doc, "3.4 Assumptions required for the reduction", 2)
    bullets(
        doc,
        [
            "Discrete binary hidden state Z ∈ {L, R}.",
            "Finite observation alphabet Y ∈ {L, R}.",
            "One-step sampling policy: acquire a cue, then make a terminal commitment.",
            "Fixed outcome preferences; no preference learning during the decision.",
            "Bayesian posterior updating under a known observation channel.",
            "Epistemic term expressed as expected information gain I(Z; Y).",
            "Pragmatic term identified with expected instrumental reward when log-preferences are proportional to reward.",
            "No parameter-learning / novelty term in the confirmatory choice contrast.",
            "Terminal reward outcomes are not additionally counted as an epistemic signal; the designed cue is the information-bearing observation of interest.",
            "Policy comparison uses the reduced expected score rather than a full process-level implementation of belief updating and neuronal dynamics.",
            "The expected-evidence-bound / ambiguity term is treated as negligible or constant under the idealized model-matching assumptions used for this discrimination exercise.",
        ],
    )
    callout(
        doc,
        "Scope of the reduced model",
        "The reduced model tests an EFE-derived epistemic-value component. It does not test the entirety of Active Inference, nor every formulation or process-level implementation of expected free energy.",
    )

    heading(doc, "4. Competing information-seeking mechanisms", 1)
    para(
        doc,
        "The candidate contribution must be positioned within a literature in which information gain is not "
        "unique to Active Inference. The main competitor is no longer “RL versus AIF” in the abstract, but "
        "alternative information-seeking objectives that can generate similar sampling.",
    )
    heading(doc, "4.1 Relation to information-theoretic exploration and experimental design", 2)
    add_table(
        doc,
        ["Framework", "Information-seeking principle", "Role in this project"],
        [
            [
                "Active Inference / EFE",
                "Expected epistemic value / information gain",
                "Source of the reduced exact-MI target under Section 3 assumptions.",
            ],
            [
                "Bayesian experimental design",
                "Expected information gain as a design utility [11,12]",
                "Shows that exact MI has a broader non-AIF normative role.",
            ],
            [
                "Intrinsic motivation / curiosity",
                "Prediction improvement, novelty, competence, or other intrinsic rewards [3]",
                "Shows that exploration is not theory-specific.",
            ],
            [
                "Information-gain exploration",
                "Policies scored by expected reduction in uncertainty",
                "Closest generic cousin of the reduced controller.",
            ],
            [
                "Information-directed sampling",
                "Mutual information used directly in exploration/exploitation [4]",
                "Strong conceptual rival to any claim that MI sensitivity identifies AIF.",
            ],
            [
                "Bayes-adaptive RL",
                "Planning in a belief MDP; exploration as value of information",
                "v3 already used a flexible reward-based member of this family as a rival.",
            ],
        ],
        col_widths=[2600, 3400, 3360],
        caption="Table 2. Neighboring information-seeking traditions. None is claimed to be exhausted by the current synthetic comparisons.",
    )
    para(
        doc,
        "Version 1.1 therefore makes no “first”, “breakthrough”, or universal novelty claim. A close human "
        "precedent already exists. Nelson et al. (2010, Experiment 3, Condition 1) constructed two binary "
        "features with the same prior and the same probability gain (0.25) but different information gain, "
        "and tested them in humans [16]. An independent recalculation recovers equal Bayes-optimal accuracy "
        "(0.75) and unequal mutual information (approximately 0.2158 vs 0.1308 nats). That is the same "
        "design class as v6 — same prior, same instrumental classification value, different exact MI — "
        "though not the same task. Design-class novelty of v6 is therefore not supported. What remains open "
        "is whether an Active Inference decision framing adds an empirically distinct prediction beyond "
        "this older optimal-experimental-design contrast (Section 11).",
    )
    heading(doc, "4.2 Programme-level rival hypotheses", 2)
    para(
        doc,
        "A further, programme-level alternative is that any engineering or behavioral success of Active "
        "Inference is explained by its information-seeking objective family without the physical or "
        "metaphysical claims sometimes attached to the Free Energy Principle — for example as an "
        "intrinsic-motivation or Bayes-adaptive control scheme [3,4]. That hypothesis is why the rival "
        "set cannot stop at generic curiosity. Information-directed sampling, Bayes-adaptive RL, Thompson "
        "sampling, UCB, Bayesian surprise, and mixture heuristics are not optional extras; they are the "
        "controls that keep an AIF-labeled fit from being re-read as support for the principle.",
    )
    add_table(
        doc,
        ["Hypothesis", "Content", "Status in this note"],
        [
            [
                "Exact-MI valuation",
                "Sampling tracks I(Z;Y) after coarse summaries are matched.",
                "Synthetically identifiable vs the tested curiosity family; human test open; weakened as a first-discovery claim by Nelson et al. [16].",
            ],
            [
                "Probability-gain valuation",
                "Sampling tracks expected increase in correct-classification probability.",
                "Supported in a close OED human cell [16]; not yet implemented as a v6 rival in our code.",
            ],
            [
                "Engineering-without-FEP",
                "AIF successes are explained as RL / intrinsic-motivation analogues.",
                "Not tested here; it is the reason the rival set must keep growing.",
            ],
            [
                "FEP-as-principle",
                "A positive epistemic-value result confirms the general Free Energy Principle.",
                "Out of scope and already blocked by the Section 1.1 ladder and by v7.",
            ],
        ],
        col_widths=[2200, 3400, 3760],
        caption="Table 2b. Programme-level hypotheses. None of these rows is new v6 evidence.",
    )

    heading(doc, "5. Adversarial sequence v1–v5", 1)
    add_figure(
        doc,
        FIG / "fig1_sequence.png",
        "Figure 1. Adversarial sequence. Each later rival narrows the conclusion that a positive result may license.",
        width=Inches(6.3),
    )

    heading(doc, "5.1 v1–v3: progressively stronger reward-based rivals", 2)
    para(
        doc,
        "The base task is a one-step hidden-context decision. A binary hidden state Z ∈ {L, R} is drawn from "
        "prior P(Z = L) = p. The agent may take a safe option, commit immediately to L or R, or pay a cost c "
        "to observe a cue and then commit. In v1–v3 the cue is a symmetric channel with reliability q = P(Y = Z).",
    )
    para(
        doc,
        "v1 separated the reduced exact-MI controller from a reward-only Bayesian controller in selected "
        "conditions. v2 then introduced parameter recovery and revealed weak-effect regions where the generating "
        "family was not reliably identifiable. v3 strengthened the reward-based rival by fitting decision "
        "precision, subjective cue reliability, and subjective information cost — a flexible Bayes-adaptive "
        "instrumental rival. Strong epistemic effects remained recoverable; weak effects did not.",
    )
    add_figure(
        doc,
        FIG / "image2.png",
        "Figure 2. v3 recovery against a flexible Bayes-adaptive reward rival. Strong epistemic weights survive "
        "the tested instrumental rival; weak weights remain an identifiability problem. These are synthetic "
        "recovery rates, not a human sample-size calculation.",
        width=Inches(5.8),
    )

    heading(doc, "5.2 v4: generic curiosity destroys specificity", 2)
    para(
        doc,
        "v4 is the crucial negative result and must remain visible. Once the rival was allowed an intrinsic "
        "information-seeking bonus of a different functional form, ordinary symmetric-cue tasks became nearly "
        "observationally equivalent. The tested coarse curiosity family scored sampling as",
    )
    equation(doc, "V_curiosity(SAMPLE)  =  V_instr  +  κ · H(Z) · (2q − 1)^γ", "5")
    para(
        doc,
        "In the tested grid, the median minimum Jensen–Shannon divergence was approximately 5.5×10⁻⁷, the "
        "90th percentile approximately 5.9×10⁻⁶, and the maximum approximately 2×10⁻⁵. A flexible coarse "
        "curiosity family can therefore closely mimic the reduced exact-MI model in the symmetric-cue task. "
        "Information seeking, as such, is not mechanism-specific.",
    )
    add_figure(
        doc,
        FIG / "fig3_v4.png",
        "Figure 3. v4 equivalence diagnostic. Minimum Jensen–Shannon divergence between the reduced exact-MI "
        "controller and the best-fitting tested curiosity parameterization, across the symmetric-cue grid.",
        width=Inches(5.6),
    )

    heading(doc, "5.3 v5: unsuccessful structural search in the symmetric task", 2)
    para(
        doc,
        "v5 is the unsuccessful structural search and should remain visible as such. A search over 175,821 "
        "pairs of symmetric-cue conditions sought robust ordering reversals between the exact-MI and curiosity "
        "families. No pair produced a 100% sign reversal across the full tested curiosity parameter grid; the "
        "best pair reversed the ordering for approximately 71.8% of rival parameterizations. This motivated a "
        "change in task geometry rather than further parameter tuning.",
    )

    heading(doc, "6. Matched asymmetric-channel design (v6)", 1)
    para(
        doc,
        "We propose a matched asymmetric-channel design. The observation channel is expanded from a single "
        "symmetric reliability q to sensitivity a = P(Y = L | Z = L) and specificity b = P(Y = R | Z = R). "
        "This makes it possible to hold coarse summaries fixed while changing the full channel geometry.",
    )
    para(doc, "Selected matched pairs satisfy:", first_indent=False)
    bullets(
        doc,
        [
            "the same prior P(Z = L);",
            "the same mean cue accuracy q̄;",
            "the same instrumental value of sampling, achieved by cost matching;",
            "a material difference in exact mutual information I(Z; Y).",
        ],
    )
    equation(doc, "I(Z; Y)  =  H(Y) − H(Y | Z)", "6")
    para(
        doc,
        "Under these constraints, a controller that scores sampling by exact mutual information predicts a "
        "sizable within-pair sampling contrast. A controller that scores sampling by a coarse "
        "uncertainty×accuracy bonus, using only q̄, predicts approximately none.",
    )
    add_figure(
        doc,
        FIG / "fig4_v6_contrast.png",
        "Figure 4. Matched-pair structural contrast for the four candidate confirmatory pairs in Table 3. "
        "Under the tested coarse curiosity model the pair is approximately indistinguishable; the exact-MI "
        "model predicts a sizable sampling contrast.",
        width=Inches(5.8),
    )

    heading(doc, "7. Structural identifiability results", 1)
    para(
        doc,
        "Table 3 reports the four candidate matched pairs retained for a later confirmatory protocol. "
        "Predicted ΔP is the absolute difference in sampling probability between the two channels of a pair, "
        "under the stated synthetic parameterization. These values are model predictions, not human data.",
    )
    add_table(
        doc,
        [
            "Pair",
            "Prior p",
            "Mean accuracy",
            "ΔMI (nats)",
            "Pred. ΔP exact-MI",
            "Pred. ΔP coarse curiosity",
        ],
        [
            ["1", "0.70", "0.780", "0.143", "0.277", "0"],
            ["2", "0.80", "0.846", "0.138", "0.270", "2.22×10⁻¹⁶"],
            ["3", "0.70", "0.738", "0.140", "0.259", "0"],
            ["4", "0.80", "0.798", "0.136", "0.245", "0"],
        ],
        col_widths=[1000, 1400, 1700, 1600, 1830, 1830],
        caption="Table 3. Candidate matched-channel conditions. ΔMI and predicted sampling contrasts are synthetic.",
    )
    para(
        doc,
        "Table 4 reports synthetic family recovery in the tested parameterization. At 240 trials, correct "
        "recovery was approximately 98.3% for exact-MI/AIF-generated data and 96.7% for the tested "
        "coarse-curiosity-generated data.",
    )
    add_table(
        doc,
        ["Trials", "True exact-MI / reduced AIF family", "True coarse-curiosity family"],
        [
            ["60", "97.2%", "80.0%"],
            ["120", "96.1%", "89.6%"],
            ["240", "98.3%", "96.7%"],
            ["480", "99.4%", "100.0%"],
        ],
        col_widths=[1800, 3780, 3780],
        caption="Table 4. v6 family recovery in the tested synthetic parameterization. Toy recovery only.",
    )
    add_figure(
        doc,
        FIG / "fig5_v6_recovery.png",
        "Figure 5. v6 family recovery in the tested synthetic parameterization. These percentages must not be "
        "used to determine participant sample size.",
        width=Inches(5.8),
    )
    callout(
        doc,
        "Sampling-size gate",
        "The current synthetic recovery rates must not be used to determine participant sample size. They are toy/synthetic recovery values under a restricted individual-agent parameterization, not a hierarchical power surface.",
    )
    para(
        doc,
        "v6 discriminates the tested accuracy-based curiosity family. It does not discriminate all conceivable "
        "information-seeking algorithms. That restriction is essential to the claim that can be defended after "
        "a later human study.",
    )

    heading(doc, "8. Equivalence and interpretation boundary (v7)", 1)
    para(
        doc,
        "v7 is not a standalone theoretical result and is not offered as a mathematically surprising theorem. "
        "Its role is to delimit the interpretation of any positive behavioral finding.",
    )
    para(
        doc,
        "Write the reduced AIF controller and an alternative information-seeking controller as",
    )
    equation(doc, "V_AIF  =  V_instr  +  α · I(Z; Y)", "7")
    equation(doc, "V_ALT  =  V_instr  +  λ · I(Z; Y)", "8")
    para(
        doc,
        "If α = λ and both controllers use the same policy mapping Q(a | x) — in the simulations, the same "
        "softmax with the same decision precision — then for every action a and every task condition x,",
    )
    equation(doc, "Q_AIF(a | x)  =  Q_ALT(a | x)", "9")
    para(
        doc,
        "The numerical grid check returned zero difference to floating-point precision. Consequently,",
    )
    equation(doc, "behavioral data   ⇏   unique theoretical derivation", "10")
    para(
        doc,
        "The equivalence result is not intended as a mathematically surprising theorem; its role is to delimit "
        "the interpretation of any positive behavioral finding. A later human contrast, even if clean, can "
        "support sensitivity to an exact-MI objective relative to specified alternatives. It cannot uniquely "
        "identify Active Inference when another theory implements the same objective and the same "
        "observation-to-action map.",
    )
    callout(
        doc,
        "Interpretation boundary",
        "A positive behavioral result can support exact-MI-sensitive valuation relative to specified alternatives. It cannot uniquely identify Active Inference or the Free Energy Principle.",
    )

    heading(doc, "9. Candidate empirical protocol (v8)", 1)
    para(
        doc,
        "Version 1.1 changes the status of the human study from a preregistration-ready empirical protocol "
        "to a preregistration-oriented candidate protocol. The protocol is not final until the following "
        "gates have been completed:",
        first_indent=False,
    )
    equation(
        doc,
        "hierarchical model  →  parameter recovery  →  false-positive analysis  →  power surface",
        "11",
    )
    callout(
        doc,
        "Frozen candidate empirical claim",
        "When instrumental value and mean cue accuracy are matched, sampling choices track exact channel mutual information rather than a generic uncertainty × accuracy proxy.",
    )

    heading(doc, "9.1 Current primary competitors", 2)
    bullets(
        doc,
        [
            "Exact-MI epistemic model: instrumental value plus a weight on I(Z; Y).",
            "Instrumental-only Bayesian controller.",
            "Coarse uncertainty × accuracy curiosity model used in v4–v6.",
        ],
    )
    para(
        doc,
        "These three models do not exhaust the relevant alternatives. The confirmatory comparison currently "
        "excludes, until implemented, at least Thompson sampling, upper-confidence-bound (UCB) rules, "
        "information-directed sampling, Bayesian-surprise variants, heuristic cue-sampling strategies, and "
        "mixture strategies. v6 discriminates the tested accuracy-based curiosity family; it does not "
        "discriminate all conceivable information-seeking algorithms.",
    )

    heading(doc, "9.2 Candidate decision rules", 2)
    add_table(
        doc,
        ["Status", "Candidate rule"],
        [
            [
                "PASS",
                "Exact-MI model outpredicts both current primary rivals on held-out confirmatory trials; ≥3 of 4 matched pairs show the preregistered direction; aggregate paired contrast lies outside a simulation-defined equivalence region.",
            ],
            [
                "FAIL",
                "A current primary rival predicts held-out choices equally well or better, or the confirmatory matched-pair contrast is absent or reversed.",
            ],
            [
                "UNRESOLVED",
                "Final hierarchical recovery is inadequate, manipulation checks fail, or effects fall inside the preregistered equivalence region.",
            ],
        ],
        col_widths=[1800, 7560],
        caption="Table 5. Candidate confirmatory decision rules. Thresholds remain unfrozen until hierarchical calibration.",
    )

    heading(doc, "10. Human feasibility and hierarchical-analysis requirements", 1)
    heading(doc, "10.1 Next empirical gate: human feasibility is now the bottleneck", 2)
    para(
        doc,
        "The dominant uncertainty is no longer whether the toy models can be separated in simulation. It is "
        "whether human participants can learn and use the asymmetric contingencies well enough for the "
        "structural contrast to survive. This is why a collaborator with access to human participants is "
        "needed before any confirmatory claim.",
    )
    para(doc, "Three risks cannot be settled by further synthetic search:", first_indent=False)
    equation(doc, "R1:  participants may fail to learn asymmetric contingencies", None)
    equation(doc, "R2:  between-subject heterogeneity may attenuate the contrast", None)
    equation(doc, "R3:  simple heuristics may reproduce the signature", None)
    add_table(
        doc,
        ["Risk", "Why simulation cannot settle it", "Pilot measurement"],
        [
            [
                "R1. Channel learnability",
                "Synthetic agents know the likelihood structure by construction.",
                "Learning curves; subjective estimates of a and b; interface comprehension checks.",
            ],
            [
                "R2. Participant heterogeneity",
                "Toy recovery uses a restricted parameter grid, not a population distribution.",
                "Empirical variance in sampling, decision precision, subjective reliability, and information cost.",
            ],
            [
                "R3. Heuristic substitution",
                "Humans may use simple cue heuristics not represented in the model set.",
                "Candidate heuristic discovery via strategy probes, belief reports, and explicit heuristic rivals.",
            ],
            [
                "Usability / floor–ceiling",
                "Interface and costs can alter sampling independently of theory.",
                "Completion rate, sampling-rate floor and ceiling, attrition, usability diagnostics.",
            ],
        ],
        col_widths=[2200, 3580, 3580],
        caption="Table 6. Human-feasibility risks and the corresponding pilot measurements.",
    )
    para(
        doc,
        "The pilot should therefore estimate learnability, interface comprehension, empirical variance, "
        "subjective estimates of a and b, sampling-rate floor and ceiling, and candidate heuristics. It is a "
        "feasibility and model-development study, not a confirmatory hypothesis test. No fixed participant "
        "number is claimed here.",
    )

    heading(doc, "10.2 Planned v9: hierarchical participant model", 2)
    para(
        doc,
        "Results of a hierarchical v9 analysis are not reported here because that analysis has not been run. "
        "What can be fixed now is the planned form. Before confirmatory recruitment, the individual-agent "
        "model must be embedded in a hierarchical population model. A representative parameterization is",
    )
    equation(doc, "log β_i  ∼  N(μ_β, σ_β²)", "12")
    para(
        doc,
        "and analogously for epistemic weights and subjective channel estimates, for example",
    )
    equation(doc, "logit(α_i / α_max)  ∼  N(μ_α, σ_α²)", "13")
    equation(doc, "logit(a_i), logit(b_i)  ∼  transformed population distributions", "14")
    para(
        doc,
        "The exact transforms and priors must be chosen before confirmatory data are opened. The purpose is "
        "not merely to fit heterogeneity but to test whether the discriminative signature remains recoverable "
        "under plausible population variation, including misperception of a and b.",
    )
    para(
        doc,
        "The v9 simulation grid should estimate, for combinations of participants N and trials T:",
        first_indent=False,
    )
    bullets(
        doc,
        [
            "P(PASS | exact-MI data-generating process);",
            "P(false PASS | each rival data-generating process);",
            "P(UNRESOLVED | model and sample size);",
            "parameter-recovery bias and interval coverage;",
            "confusion matrices under heuristic and mixture misspecification.",
        ],
    )
    callout(
        doc,
        "Confirmatory calibration rule",
        "Final participant-level sample size and confirmatory decision thresholds will be set only after simulation-based calibration of this hierarchical model.",
    )

    heading(doc, "10.3 Additional observables", 2)
    para(
        doc,
        "Reaction time, confidence, pupil responses, EEG/MEG/fMRI, or belief reports could become useful only "
        "if competing process theories make preregistered quantitative predictions for those observables. The "
        "current document does not claim such predictions. In particular, a statement such as “higher mutual "
        "information should cause longer reaction time under AIF” does not follow automatically from the EFE "
        "objective. A process model connecting latent inference dynamics to those measures would be required. "
        "Version 1.1 therefore treats them as future routes around the v7 equivalence ceiling, not as current evidence.",
    )

    heading(doc, "11. Prior art and scope of novelty", 1)
    para(
        doc,
        "The proposed matched-channel construction may be methodologically useful, but design-class novelty "
        "is not supported. Nelson (2005) already searched for environments in which information-gain, "
        "probability-gain, diagnosticity, and impact diverge [5]. Nelson et al. (2010) then used computer "
        "optimization to build such environments and ran them in humans [16]. Their Experiment 3, Condition 1 "
        "ties probability gain while varying information gain — the same structural idea as v6. After "
        "experience-based learning, only 12 of 22 participants (55%) preferred the higher-information-gain "
        "feature. That result is a small-N human cell, not a decisive null, but it is a relevant anti-context "
        "for any claim that exact-MI preference will be large once classification value is matched. Bayesian "
        "experimental design had already treated expected information gain as a design utility [11,12].",
    )
    callout(
        doc,
        "Current novelty status: design-class claim not supported",
        "Nelson et al. (2010) already tested a same-prior, same-probability-gain, different-information-gain contrast in humans. v6 is an AIF-framed analogue, not a historically first discriminator.",
    )
    para(
        doc,
        "A dedicated novelty audit should search at least the following literatures:",
        first_indent=False,
    )
    bullets(
        doc,
        [
            "Bayesian experimental design and optimal sensing;",
            "human question selection, diagnosticity, probability gain, and information gain;",
            "intrinsic motivation and curiosity;",
            "information-directed sampling and information-theoretic reinforcement learning;",
            "Bayes-adaptive RL and value-of-information experiments;",
            "asymmetric binary observation channels and matched diagnostic tests.",
        ],
    )

    heading(doc, "12. Limitations", 1)
    bullets(
        doc,
        [
            "The current exact-MI controller is a reduced policy-ranking special case, not a full process-level Active Inference implementation.",
            "All reported v1–v8 results are synthetic; no human empirical data have been collected.",
            "The v6 comparison is against specified rival families and does not exclude all possible information-seeking algorithms.",
            "Sequential learning, temporal credit assignment, and model learning are deliberately absent from the one-step task.",
            "The strongest positive result is objective-level discrimination, not theory-level identification.",
            "Human participants may mislearn or simplify the asymmetric channels, which could erase the predicted contrast.",
            "The current recovery results are not a substitute for a hierarchical power and false-positive analysis.",
            "RT, confidence, and neural measures require separate process-level predictions before they can be used as discriminators.",
        ],
    )
    add_table(
        doc,
        ["Supported synthetically", "Not established"],
        [
            [
                "Reward-only rivals can be too weak to assess epistemic specificity.",
                "Humans implement exact mutual-information valuation.",
            ],
            [
                "A tested generic curiosity family can be nearly observationally equivalent in symmetric tasks.",
                "The v6 design is historically first.",
            ],
            [
                "Matched asymmetric channels can synthetically separate exact-MI sensitivity from the tested coarse-curiosity family.",
                "A positive v6 result would uniquely validate Active Inference or FEP.",
            ],
            [
                "A same-objective alternative is behaviorally equivalent under the same policy map.",
                "The final confirmatory sample size is known.",
            ],
        ],
        col_widths=[4680, 4680],
        caption="Table 7. What the present evidence does and does not establish.",
    )

    heading(doc, "13. Discussion", 1)
    heading(doc, "13.1 Level of inference", 2)
    para(
        doc,
        "The adversarial sequence changes the research question from “Does Active Inference fit exploratory "
        "behavior?” to “Which information-sensitive objective, if any, is required to explain behavior under "
        "a task designed to separate coarse accuracy from full channel information geometry?” This is a "
        "narrower question, but it is also more falsifiable. It occupies the identifiable-prediction and "
        "strong-rival rungs of the Section 1.1 ladder. It does not move a result from a specific EFE model "
        "back up to the Free Energy Principle.",
    )
    para(
        doc,
        "The main candidate methodological contribution is the matched asymmetric-channel manipulation. "
        "Rather than search for extreme parameters within a structurally non-identifying symmetric task, the "
        "design holds coarse summaries constant and varies the property of interest — exact mutual information. "
        "The main conceptual contribution is the interpretation boundary: even an exact-MI-sensitive behavioral "
        "signature identifies an objective only relative to alternatives; it does not identify a unique "
        "theoretical derivation.",
    )
    para(
        doc,
        "The current rival set is still too small to support a sweeping claim. Instrumental Bayes plus one "
        "curiosity proxy do not exhaust the alternatives. At minimum, a later analysis should treat as open "
        "rivals Thompson sampling [13], UCB [14], information-directed sampling [4], Bayesian surprise [15], "
        "heuristic cue sampling, and mixture strategies. Until those families are implemented, the honest "
        "statement is that v6 discriminates the tested accuracy-based curiosity family; it does not discriminate "
        "all conceivable information-seeking algorithms.",
    )
    para(
        doc,
        "The appropriate next step for this paper is therefore not another unbounded synthetic search. It is "
        "a sequence of gates already named above: formal mapping to Nelson et al. (2010), attempted reanalysis "
        "of existing human OED data, hierarchical v9 calibration, and only then an AIF-framed feasibility pilot.",
    )
    heading(doc, "13.2 Place in a broader research programme", 2)
    para(
        doc,
        "Epistemic-value identifiability is one kill-test, not the whole FEP/AIF programme. Neighboring "
        "tests — a preregistered closed-loop intervention comparison against named control algorithms, a "
        "coarse-graining / scale-gap analysis, independent clinical replication, and process-level "
        "discrimination — answer different questions and are not licensed by a v6 result. This note does "
        "not claim those tests, and the strategic FEP-audit used to draw the map is not evidence for the "
        "synthetic v6 contrasts.",
    )
    add_table(
        doc,
        ["Programme slot", "Question", "Role of this note"],
        [
            [
                "Epistemic-value identifiability",
                "Can exact information gain be distinguished from specified information-seeking alternatives?",
                "This paper: synthetic v1–v7; human feasibility pending.",
            ],
            [
                "Closed-loop intervention",
                "Does a preregistered generative AIF model beat named controls on new interventions?",
                "Out of scope. Different estimand.",
            ],
            [
                "Coarse-graining / scale-gap",
                "Do variational or Markov-blanket properties survive changes of scale?",
                "Out of scope.",
            ],
            [
                "Clinical replication",
                "Do information-seeking or precision accounts replicate in independent cohorts?",
                "Out of scope.",
            ],
            [
                "Process-level discrimination",
                "Do RT, confidence, or neural measures split AIF process theory from same-objective alternatives?",
                "Named as a future route around v7; no predictions claimed.",
            ],
        ],
        col_widths=[2400, 3600, 3360],
        caption="Table 8. Research programme. Rows other than the first are positioning, not results.",
    )

    heading(doc, "14. Conclusion", 1)
    para(
        doc,
        "Version 1.1 reframes the project as a methodological proposal with synthetic identifiability evidence "
        "and pending human validation. The reduced exact-MI controller is now explicitly tied to an EFE "
        "decomposition under stated assumptions; its limits are equally explicit. Generic information seeking "
        "is not unique to AIF, and a same-objective alternative creates an unavoidable behavioral equivalence ceiling.",
    )
    para(
        doc,
        "The proposed human target is therefore modest but testable: determine whether sampling choices are "
        "sensitive to exact observation-channel mutual information after matching prior, mean accuracy, and "
        "instrumental value, relative to specified instrumental and coarse-curiosity alternatives. If this "
        "contrast survives hierarchical calibration and human feasibility testing, it becomes a defensible "
        "preregistered experiment. If it does not, the design should be revised rather than the claim rescued post hoc.",
    )
    callout(
        doc,
        "Current status",
        "METHOD PROPOSED + SYNTHETICALLY IDENTIFIABLE + HUMAN VALIDATION PENDING. This is not a preregistration-ready confirmatory protocol.",
    )

    heading(doc, "Appendix A. Full equations", 1)
    para(
        doc,
        "Hidden state and actions. Z ∈ {L, R} with P(Z = L) = p. Actions at stage 0:",
        first_indent=False,
    )
    bullets(
        doc,
        [
            "SAFE → deterministic reward r_safe;",
            "COMMIT_L → r_win if Z = L else r_lose;",
            "COMMIT_R → r_win if Z = R else r_lose;",
            "SAMPLE → pay cost c, observe cue Y, then commit optimally.",
        ],
    )
    para(doc, "Symmetric channel (v1–v5). P(Y = Z) = q, and", first_indent=False)
    equation(doc, "P(Y = L)  =  p q  +  (1 − p)(1 − q)", "A1")
    equation(doc, "I(Z; Y)  =  H(P(Y = L))  −  H(q)", "A2")
    para(
        doc,
        "Asymmetric channel (v6). Sensitivity a = P(Y = L | Z = L), specificity b = P(Y = R | Z = R), and",
        first_indent=False,
    )
    equation(doc, "P(Y = L)  =  p a  +  (1 − p)(1 − b)", "A3")
    equation(doc, "H(Y | Z)  =  p H(a)  +  (1 − p) H(b)", "A4")
    equation(doc, "I(Z; Y)  =  H(P(Y = L))  −  H(Y | Z)", "A5")
    para(
        doc,
        "Instrumental value of sampling. After observing Y the agent commits to the posterior-optimal side:",
        first_indent=False,
    )
    equation(doc, "V_instr(SAMPLE)  =  −c  +  Σ_y P(y) max{V(COMMIT_L | y), V(COMMIT_R | y)}", "A6")
    para(doc, "Reduced AIF score and policy mapping.", first_indent=False)
    equation(doc, "V_AIF(SAMPLE)  =  V_instr(SAMPLE)  +  α_epi I(Z; Y)", "A7")
    equation(doc, "Q(a | x)  =  exp(β V_a(x))  /  Σ_{a′} exp(β V_{a′}(x))", "A8")
    para(doc, "Tested coarse curiosity family.", first_indent=False)
    equation(doc, "V_curiosity(SAMPLE)  =  V_instr(SAMPLE)  +  κ H(Z) (2 q̄ − 1)^γ", "A9")
    para(
        doc,
        "Binary entropy uses the natural logarithm, so all information quantities are in nats: "
        "H(u) = −u log u − (1 − u) log(1 − u), with H(0) = H(1) = 0.",
        first_indent=False,
    )

    heading(doc, "Appendix B. Verification note: units of mutual information", 1)
    para(
        doc,
        "All entropies in the simulation code use natural logarithms and are therefore reported in nats. "
        "For the previously discussed symmetric example p = 0.64 and q = 0.95:",
        first_indent=False,
    )
    equation(doc, "P(Y = L)  =  0.64 × 0.95  +  0.36 × 0.05  =  0.626", "B1")
    equation(doc, "H(Y)  =  H(0.626)  ≈  0.6610 nats", "B2")
    equation(doc, "H(Y | Z)  =  H(0.95)  ≈  0.1985 nats", "B3")
    equation(doc, "I(Z; Y)  ≈  0.4625 nats", "B4")
    para(
        doc,
        "A value near 0.953 for H(0.626) corresponds to bits, not nats. Mixing the two units creates a false "
        "discrepancy. The implemented CSV value of approximately 0.4625 nats is consistent with the formula "
        "above. Entropies are reported in nats (natural logarithms).",
        first_indent=False,
    )

    heading(doc, "Appendix C. Simulation parameters", 1)
    para(
        doc,
        "The following values are those used for the reported synthetic claims. They are recorded here so that "
        "an independent reimplementation can target the same grid. They are not human-protocol parameters.",
        first_indent=False,
    )
    add_table(
        doc,
        ["Quantity", "Value used in reported simulations"],
        [
            ["Actions", "SAFE, COMMIT_L, COMMIT_R, SAMPLE"],
            ["Default reward scale", "r_win = 1, r_lose = 0, r_safe = 0.55"],
            ["Default decision precision", "β = 8"],
            ["Default epistemic weight (v1 grid)", "α_epi = 1"],
            [
                "v1 search grid (from aif_kill_test.py)",
                "p ∈ [0.50, 0.95], q ∈ [0.55, 0.95], c ∈ [0.00, 0.40]",
            ],
            ["Information units", "nats (natural log)"],
            ["v5 pair search size", "175,821 symmetric-cue pairs"],
            [
                "v5 best reversal coverage",
                "approximately 71.8% of tested curiosity parameterizations",
            ],
            ["v6 recovery trial counts", "60, 120, 240, 480 synthetic trials"],
            [
                "v7 numerical check",
                "zero difference to floating-point precision under α = λ and identical softmax",
            ],
        ],
        col_widths=[3600, 5760],
        caption="Table C1. Recorded simulation parameters for the synthetic claims. Hierarchical v9 priors are not yet fixed.",
    )

    heading(doc, "Appendix D. Reproducibility artifacts and current state", 1)
    bullets(
        doc,
        [
            "aif_kill_test.py",
            "v3_alpha_recovery.csv",
            "v4_equivalence_stats.csv",
            "v5_all_structural_pairs.csv",
            "v6_selected_channel_pairs.csv",
            "v6_structural_checks.csv",
            "v6_family_recovery.csv",
            "v7_equivalence_summary.csv",
            "v8_preregistration_protocol.json",
            "state_checkpoint_v8.json",
        ],
    )
    para(
        doc,
        "These artifacts support the synthetic claims reported here. Independent reimplementation has not yet "
        "been completed and remains a required gate before a confirmatory human study.",
        first_indent=False,
    )

    heading(doc, "Appendix E. Changes from Version 1.0", 1)
    bullets(
        doc,
        [
            "Removed AIF-KILL from the public-facing title; retained it only as an internal project identifier.",
            "Added a formal relation between the reduced score and the expected-free-energy decomposition, with explicit assumptions and an explicit status for α_epi.",
            "Identified Nelson et al. 2010 Experiment 3 Condition 1 as a close human OED precedent; design-class novelty of v6 is not supported.",
            "Added a subsection relating the design to information-theoretic exploration and experimental design.",
            "Recast v7 as an interpretation/identifiability boundary rather than a standalone theoretical result, with the formal Q1 = Q2 statement.",
            "Changed v8 status from “preregistration-ready” to “preregistration-oriented candidate protocol”.",
            "Made the synthetic-recovery sampling-size prohibition a central gate.",
            "Added a Next empirical gate section with human-feasibility risks R1–R3 and explicit pilot tasks.",
            "Added planned hierarchical v9 as an analysis plan only, with no invented results.",
            "Expanded the rival set to Thompson sampling, UCB, information-directed sampling, Bayesian surprise, heuristic cue sampling, and mixture strategies.",
            "Added a units verification note for mutual information (nats versus bits).",
            "Clarified that RT, confidence, and neural measures require separate process-level predictions before they can be used as discriminators.",
            "Changed the external status from PREREGISTRATION READY to METHOD PROPOSED + SYNTHETICALLY IDENTIFIABLE + HUMAN VALIDATION PENDING.",
            "Added a FEP→model→identifiable-prediction→rival motivation ladder, an explicit scope box, programme-level rival hypotheses, and a research-programme map. These are positioning, not new v6 evidence.",
        ],
    )

    heading(doc, "References", 1)
    refs = [
        "[1] Friston K, Rigoli F, Ognibene D, Mathys C, Fitzgerald T, Pezzulo G. Active inference and epistemic value. Cognitive Neuroscience. 2015;6(4):187–214. doi:10.1080/17588928.2015.1020053.",
        "[2] Da Costa L, Parr T, Sajid N, Veselic S, Neacsu V, Friston K. Active inference on discrete state-spaces: A synthesis. Journal of Mathematical Psychology. 2020;99:102447. doi:10.1016/j.jmp.2020.102447.",
        "[3] Oudeyer P-Y, Kaplan F. What is intrinsic motivation? A typology of computational approaches. Frontiers in Neurorobotics. 2007;1:6. doi:10.3389/neuro.12.006.2007.",
        "[4] Russo D, Van Roy B. Learning to optimize via information-directed sampling. Advances in Neural Information Processing Systems 27. 2014.",
        "[5] Nelson JD. Finding useful questions: on Bayesian diagnosticity, probability, impact, and information gain. Psychological Review. 2005;112(4):979–999. doi:10.1037/0033-295X.112.4.979.",
        "[6] Chou K-P, Hakimi N, Hsu T-Y, Smith R. A systematic empirical comparison of active inference and reinforcement learning models in accounting for decision-making under uncertainty. Journal of Mathematical Psychology. 2026;130:103011. doi:10.1016/j.jmp.2026.103011.",
        "[7] Lageman J, Fahrenfort JJ, Slagter HA. Prediction in action: Toward an empirical science of active inference. Neuroscience & Biobehavioral Reviews. 2026;188:106817. doi:10.1016/j.neubiorev.2026.106817.",
        "[8] Van Wallendael LR. Implicit diagnosticity in an information-buying task. Journal of Behavioral Decision Making. 1995;8:245–264. doi:10.1002/bdm.3960080403.",
        "[9] Skov RB, Sherman SJ. Information-gathering processes: Diagnosticity, hypothesis-confirmatory strategies, and perceived hypothesis confirmation. Journal of Experimental Social Psychology. 1986;22(2):93–121. doi:10.1016/0022-1031(86)90031-4.",
        "[10] Friston K. Action and behavior: a free-energy formulation. Biological Cybernetics. 2010;102:227–260. doi:10.1007/s00422-010-0364-z.",
        "[11] Lindley DV. On a measure of the information provided by an experiment. Annals of Mathematical Statistics. 1956;27(4):986–1005.",
        "[12] Chaloner K, Verdinelli I. Bayesian experimental design: A review. Statistical Science. 1995;10(3):273–304.",
        "[13] Thompson WR. On the likelihood that one unknown probability exceeds another in view of the evidence of two samples. Biometrika. 1933;25(3/4):285–294.",
        "[14] Auer P, Cesa-Bianchi N, Fischer P. Finite-time analysis of the multiarmed bandit problem. Machine Learning. 2002;47:235–256.",
        "[15] Itti L, Baldi P. Bayesian surprise attracts human attention. Vision Research. 2009;49(10):1295–1306. doi:10.1016/j.visres.2008.09.007.",
        "[16] Nelson JD, McKenzie CRM, Cottrell GW, Sejnowski TJ. Experience matters: Information acquisition optimizes probability gain. Psychological Science. 2010;21(7):960–969. doi:10.1177/0956797610372637.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(ref)
        set_run_font(r, size=10)

    doc.core_properties.title = "Structural Identifiability of Epistemic Value in Active Inference"
    doc.core_properties.subject = "Version 1.1 research note — METHOD PROPOSED"
    doc.core_properties.comments = (
        "Version 1.1 outreach revision. Internal ID AIF-KILL-001. "
        "No new experimental results. Status: METHOD PROPOSED + SYNTHETICALLY IDENTIFIABLE + HUMAN VALIDATION PENDING."
    )
    doc.save(OUT)
    print("wrote", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    build()
